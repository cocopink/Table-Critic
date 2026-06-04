"""
导出 WikiTQ / TabFact bad case 分析到 Excel。
直接复用项目自带的 evaluate 函数做准确率判断。
"""

import pickle
import sys
import os
import importlib.util


def _load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


qa_evaluate = _load_module('qa_evaluate', 'refine/TableQA/utils/evaluate.py')
fv_evaluate = _load_module('fv_evaluate', 'refine/TableFV/utils/evaluate.py')


def extract_chain_info(sample):
    chain = sample.get('chain', [])
    operations = [step.get('operation_name', '') for step in chain]
    thoughts = []
    for step in chain:
        t = step.get('thought', '')
        if t:
            thoughts.append(t[:120])
    return operations, thoughts


def get_pred_from_chain(sample):
    try:
        results = sample["chain"][-1]["parameter_and_conf"]
        return results[0][0]
    except (KeyError, IndexError):
        return 'ERROR: no prediction'


def export_dataset(name, pkl_path, out_path, eval_one_fn, get_correct_fn):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill

    print(f"\n{'='*50}")
    print(f"[{name}] Loading: {pkl_path}")
    with open(pkl_path, 'rb') as f:
        data = pickle.load(f)

    correct, wrong = [], []
    for s in data:
        try:
            is_correct, pred, correct_ans = eval_one_fn(s)
        except Exception as e:
            is_correct, pred, correct_ans = None, 'ERROR', str(e)

        if is_correct is True:
            correct.append((s, pred, correct_ans))
        else:
            wrong.append((s, pred, correct_ans))

    total = len(correct) + len(wrong)
    acc = len(correct) / total * 100 if total > 0 else 0
    print(f"[{name}] Total: {total}, Correct: {len(correct)}, Wrong: {len(wrong)}, Accuracy: {acc:.2f}%")

    # Build Excel
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{name} Bad Cases"

    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='2F5496', end_color='2F5496', fill_type='solid')

    headers = ['ID', 'Question', 'Pred Answer', 'Correct Answer',
               'Table Caption', 'Table Text',
               'Operations', 'Chain Length', 'Thought Summary', 'Judge']
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')

    for row_idx, (s, pred, correct_ans) in enumerate(wrong, 2):
        operations, thoughts = extract_chain_info(s)
        ws.cell(row=row_idx, column=1, value=s.get('ids', s.get('id', '')))
        ws.cell(row=row_idx, column=2, value=s.get('statement', ''))
        ws.cell(row=row_idx, column=3, value=str(pred))
        ws.cell(row=row_idx, column=4, value=str(correct_ans))
        ws.cell(row=row_idx, column=5, value=s.get('table_caption', ''))
        ws.cell(row=row_idx, column=6, value=str(s.get('table_text', '')))
        ws.cell(row=row_idx, column=7, value=' -> '.join(operations) if operations else 'N/A')
        ws.cell(row=row_idx, column=8, value=len(operations))
        ws.cell(row=row_idx, column=9, value=' | '.join(thoughts))
        ws.cell(row=row_idx, column=10, value=str(s.get('judge', '')))

    for col_letter, w in [('A', 14), ('B', 60), ('C', 20), ('D', 20),
                          ('E', 30), ('F', 60), ('G', 40), ('H', 12), ('I', 80), ('J', 16)]:
        ws.column_dimensions[col_letter].width = w
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical='top')

    # Statistics sheet
    ws2 = wb.create_sheet("Statistics")
    ws2.cell(row=1, column=1, value='Metric').font = header_font
    ws2.cell(row=1, column=2, value='Value').font = header_font
    ws2.cell(row=1, column=1).fill = header_fill
    ws2.cell(row=1, column=2).fill = header_fill

    op_counts = {}
    for s, _, _ in wrong:
        for step in s.get('chain', []):
            op = step.get('operation_name', '')
            if op:
                op_counts[op] = op_counts.get(op, 0) + 1

    stats = [
        ('Dataset', name),
        ('Total Samples', total),
        ('Correct', len(correct)),
        ('Wrong (Bad Cases)', len(wrong)),
        ('Accuracy', f"{acc:.2f}%"),
        ('', ''),
        ('Operation Distribution in Bad Cases', ''),
    ]
    for op, count in sorted(op_counts.items(), key=lambda x: -x[1]):
        stats.append((op, count))

    for row_idx, (metric, value) in enumerate(stats, 2):
        ws2.cell(row=row_idx, column=1, value=metric)
        ws2.cell(row=row_idx, column=2, value=str(value))

    wb.save(out_path)
    print(f"[{name}] Saved to {out_path}")


def wikitq_eval_one_wrapper(target_values_map):
    """Returns a function that evaluates one WikiTQ sample."""
    def eval_one(sample):
        try:
            is_correct = qa_evaluate.wikitq_match_func(sample, target_values_map[sample['ids']])
        except Exception:
            is_correct = None
        pred = get_pred_from_chain(sample)
        correct_ans = str(sample.get('answer', ''))
        return is_correct, pred, correct_ans
    return eval_one


def tabfact_eval_one_wrapper():
    """Returns a function that evaluates one TabFact sample."""
    def eval_one(sample):
        try:
            is_correct = fv_evaluate.tabfact_match_func(sample)
        except Exception:
            is_correct = None
        pred = get_pred_from_chain(sample)
        label = sample.get('label', -1)
        return is_correct, pred, label
    return eval_one


def main():
    os.makedirs('refine-logs', exist_ok=True)

    # --- WikiTQ ---
    print("[WikiTQ] Loading targets...")
    # Load target map inline (same logic as evaluate.py)
    from codecs import open as codecs_open
    tsv_unescape_list = qa_evaluate.tsv_unescape_list
    to_value_list = qa_evaluate.to_value_list
    target_values_map = {}
    tagged_dataset_path = 'thought/TableQA/data/wikitq/tagged_data'
    for filename in os.listdir(tagged_dataset_path):
        if filename[0] == '.': continue
        fn = os.path.join(tagged_dataset_path, filename)
        with codecs_open(fn, 'r', 'utf8') as fin:
            header = fin.readline().rstrip('\n').split('\t')
            for line in fin:
                stuff = dict(zip(header, line.rstrip('\n').split('\t')))
                ex_id = stuff['id']
                original_strings = tsv_unescape_list(stuff['targetValue'])
                canon_strings = tsv_unescape_list(stuff['targetCanon'])
                target_values_map[ex_id] = to_value_list(original_strings, canon_strings)
    print(f"[WikiTQ] Loaded {len(target_values_map)} targets")

    # ===== Refine Stage =====
    export_dataset(
        'WikiTQ (Refine)',
        'results/refine/wikitq/gpt-5.4/final_result.pkl',
        'refine-logs/bad_cases_wikitq_gpt54_refine.xlsx',
        wikitq_eval_one_wrapper(target_values_map),
        None,
    )

    export_dataset(
        'TabFact (Refine)',
        'results/refine/tabfact/gpt-5.4/final_result.pkl',
        'refine-logs/bad_cases_tabfact_gpt54_refine.xlsx',
        tabfact_eval_one_wrapper(),
        None,
    )

    # ===== Thought Stage =====
    export_dataset(
        'WikiTQ (Thought)',
        'results/thought/wikitq/gpt-5.4/final_result.pkl',
        'refine-logs/bad_cases_wikitq_gpt54_thought.xlsx',
        wikitq_eval_one_wrapper(target_values_map),
        None,
    )

    export_dataset(
        'TabFact (Thought)',
        'results/thought/tabfact/gpt-5.4/final_result.pkl',
        'refine-logs/bad_cases_tabfact_gpt54_thought.xlsx',
        tabfact_eval_one_wrapper(),
        None,
    )


if __name__ == '__main__':
    main()
